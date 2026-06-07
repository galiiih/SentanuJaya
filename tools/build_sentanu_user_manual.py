from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "D:/OdooProject/Panduan Operasional ERP Sentanu Jaya.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
INK = "0B2545"
BORDER = "B8C2CC"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width.inches * 1440)))


def style_cell_text(cell, bold=False, color=INK, size=9.5):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths=None, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = header
        set_cell_shading(hdr[idx], header_fill)
        set_cell_margins(hdr[idx])
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        style_cell_text(hdr[idx], bold=True, color=INK, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            style_cell_text(cells[idx], size=9.2)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)
    set_table_borders(table, color="D6DEE8", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    for run in p2.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
    return p


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
    return p


def section_break(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("PANDUAN OPERASIONAL ERP SENTANU JAYA")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    r = subtitle.add_run("Odoo 17 Community - Workflow Barang Masuk, Produksi, Pengiriman, Invoice, dan Claim")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Versi 1.0 | Disusun 7 Juni 2026 | Database: SentanuJaya")
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    add_callout(
        doc,
        "Tujuan dokumen",
        "Dokumen ini menjelaskan cara mengoperasikan ERP Sentanu Jaya dari setup awal sampai barang keluar, termasuk cara membuat user, menambah gudang, mengisi master data, memproses Surat Jalan Masuk, Production Batch, Work Order, Final QC, Surat Jalan Keluar, Invoice/Payment, Customer Claim, Claim Rework, dan Redelivery.",
    )
    add_heading(doc, "Daftar Isi Ringkas", 1)
    add_table(
        doc,
        ["Bab", "Isi"],
        [
            ["1", "Ringkasan alur utama ERP Sentanu Jaya."],
            ["2", "Persiapan awal: login, user, employee/PIC, dan gudang."],
            ["3", "Master data: customer, product, warna, service, defect, claim reason, pricelist."],
            ["4", "Workflow barang masuk sampai keluar."],
            ["5", "Workflow Customer Claim dan Redelivery."],
            ["6", "Cabang keputusan penting dan exception handling."],
            ["7", "Checklist harian per bagian."],
            ["8", "Menu referensi cepat."],
            ["9", "Troubleshooting operasional."],
            ["10", "Contoh simulasi transaksi end-to-end."],
            ["11", "Batasan versi saat ini."],
            ["12", "Rekomendasi penggunaan saat go-live."],
        ],
        widths=[Inches(0.7), Inches(5.8)],
    )

    add_heading(doc, "1. Ringkasan Alur Utama", 1)
    add_para(
        doc,
        "Alur utama Sentanu Jaya dimulai dari barang customer masuk dengan Surat Jalan Masuk, diperiksa Incoming QC, diproses melalui Production Batch dan Work Order, diperiksa Final QC, dikirim kembali dengan Surat Jalan Keluar, lalu ditagihkan melalui Invoice dan Payment. Jika setelah pengiriman ada komplain, proses dilanjutkan lewat Customer Claim, Claim Rework, dan Surat Jalan Pengiriman Ulang.",
    )
    add_table(
        doc,
        ["Urutan", "Tahap", "Menu Utama", "Output"],
        [
            ["1", "Customer dan master data", "Contacts, Sentanu Jaya/Master Data", "Customer, Item, Color, Service, Pricelist, Defect"],
            ["2", "Barang masuk", "Sentanu Jaya/Receiving/Surat Jalan Masuk", "SJM dan Incoming QC"],
            ["3", "Batch produksi", "Sentanu Jaya/Production/Production Batches", "PB dan WO Raw/Paint/Polish"],
            ["4", "Pengerjaan produksi", "Sentanu Jaya/Production/Work Orders", "Status WO done dan material usage"],
            ["5", "Final QC", "Sentanu Jaya/Production/Final QC", "Batch QC Passed atau QC Failed"],
            ["6", "Pengiriman", "Sentanu Jaya/Delivery/Surat Jalan Keluar", "SJK normal atau SJU redelivery"],
            ["7", "Penagihan", "Accounting/Customers/Invoices", "Invoice dan Payment"],
            ["8", "Claim", "Sentanu Jaya/Claims/Customer Claims", "Claim valid/invalid, Claim Rework, Redelivery"],
        ],
        widths=[Inches(0.55), Inches(1.35), Inches(2.35), Inches(2.25)],
    )

    add_heading(doc, "2. Persiapan Awal Sistem", 1)
    add_heading(doc, "2.1 Login dan Halaman Utama", 2)
    number(doc, "Buka browser dan masuk ke http://localhost:8069.")
    number(doc, "Login menggunakan akun administrator atau akun user yang sudah dibuat.")
    number(doc, "Pastikan menu Sentanu Jaya muncul di deretan aplikasi/menu. Jika belum muncul, refresh browser atau logout/login kembali.")
    add_callout(
        doc,
        "Catatan akses",
        "Modul Sentanu Jaya saat ini memakai akses internal user Odoo. Untuk operasional real, sebaiknya nanti dibuat security group khusus: Admin Sentanu, Receiving, Produksi, QC, Delivery, dan Accounting.",
    )

    add_heading(doc, "2.2 Membuat Account Pengguna", 2)
    add_para(doc, "Akun pengguna dibuat dari menu standar Odoo. Gunakan langkah ini untuk admin, staff receiving, PIC produksi, QC, delivery, dan accounting.")
    number(doc, "Masuk ke Settings.")
    number(doc, "Pilih Users & Companies, lalu Users.")
    number(doc, "Klik New.")
    number(doc, "Isi Name, Email, dan Username/Login.")
    number(doc, "Pada bagian Access Rights, aktifkan aplikasi yang sesuai: Contacts, Inventory, Manufacturing, Accounting, Employees, dan akses internal user.")
    number(doc, "Klik Save.")
    number(doc, "Jika perlu mengirim undangan login, gunakan tombol Send Invitation atau reset password dari user tersebut.")
    add_table(
        doc,
        ["Jenis User", "Hak Akses Minimum", "Keterangan"],
        [
            ["Admin ERP", "Settings, Contacts, Inventory, Manufacturing, Accounting, Sentanu Jaya", "Mengelola konfigurasi dan master data."],
            ["Admin Receiving", "Contacts, Inventory, Sentanu Jaya", "Input SJ Masuk dan Incoming QC."],
            ["Supervisor Produksi", "Manufacturing, Employees, Sentanu Jaya", "Membuat batch, WO, dan monitor progress."],
            ["QC", "Sentanu Jaya", "Input Incoming QC dan Final QC."],
            ["Delivery", "Inventory, Sentanu Jaya", "Membuat Surat Jalan Keluar dan Redelivery."],
            ["Accounting", "Accounting, Contacts, Sentanu Jaya", "Membuat invoice dan mencatat payment."],
        ],
        widths=[Inches(1.35), Inches(2.35), Inches(2.8)],
    )

    add_heading(doc, "2.3 Membuat Employee / PIC", 2)
    number(doc, "Masuk ke Employees.")
    number(doc, "Klik New.")
    number(doc, "Isi nama karyawan, department, jabatan, dan informasi kontak.")
    number(doc, "Klik Save.")
    add_para(doc, "Employee dipakai sebagai PIC di Work Order dan Final QC. Contoh PIC: Galih untuk WO Raw, Asep untuk WO Paint, Budi untuk WO Polish.")

    add_heading(doc, "2.4 Menambahkan Gudang", 2)
    add_para(doc, "Gudang dibuat dari menu Inventory standar Odoo. Modul Sentanu Jaya saat ini belum membuat stock move otomatis, tetapi setup gudang tetap penting untuk inventory material dan perkembangan integrasi berikutnya.")
    number(doc, "Masuk ke Inventory.")
    number(doc, "Pilih Configuration, lalu Warehouses.")
    number(doc, "Klik New.")
    number(doc, "Isi Warehouse Name, misalnya Sentanu Jaya Workshop.")
    number(doc, "Isi Short Name, misalnya SJ.")
    number(doc, "Klik Save.")
    number(doc, "Jika butuh lokasi detail, masuk ke Configuration > Locations, lalu buat lokasi seperti Raw Area, Paint Area, Polish Area, QC Area, Finished Goods, dan Reject Area.")
    add_table(
        doc,
        ["Lokasi Disarankan", "Fungsi"],
        [
            ["Raw Area", "Barang masuk yang menunggu atau sedang proses persiapan permukaan."],
            ["Paint Area", "Barang yang masuk proses Base Coat, Color Layer, Clear Coat, atau Oven."],
            ["Polish Area", "Barang yang masuk Wet Sanding, Compound, atau Polishing."],
            ["QC Area", "Barang yang menunggu Incoming QC atau Final QC."],
            ["Finished Goods", "Barang yang sudah lolos QC dan siap dikirim."],
            ["Reject Area", "Barang reject incoming atau gagal QC yang menunggu keputusan return/rework."],
        ],
        widths=[Inches(1.8), Inches(4.7)],
    )

    add_heading(doc, "3. Master Data yang Wajib Diisi", 1)
    add_para(doc, "Sebelum transaksi harian, isi master data berikut. Master data yang rapi akan membuat input transaksi lebih cepat dan laporan lebih akurat.")
    add_heading(doc, "3.1 Customer dan Tax Information", 2)
    number(doc, "Masuk ke Contacts.")
    number(doc, "Klik New.")
    number(doc, "Isi nama customer, alamat, nomor telepon, email, dan data lain yang tersedia.")
    number(doc, "Buka tab Tax Information.")
    number(doc, "Isi NPWP jika ada.")
    number(doc, "Pilih Tax Profile: NON, PPN, PPH, atau BOTH.")
    number(doc, "Sistem akan membantu mengisi Default PPN, Default PPh, dan PPh Rate sesuai Tax Profile.")
    number(doc, "Klik Save.")

    add_heading(doc, "3.2 Item / Product Master", 2)
    number(doc, "Masuk ke Inventory atau Sales, lalu Products.")
    number(doc, "Klik New.")
    number(doc, "Isi Product Name, misalnya Roof Garnish Mazda, Spoiler, Cover Velg.")
    number(doc, "Pilih Product Type sesuai kebutuhan. Untuk jasa murni dapat memakai Service, untuk barang/inventory material gunakan Storable Product.")
    number(doc, "Klik Save.")

    add_heading(doc, "3.3 Color Master", 2)
    number(doc, "Masuk ke Sentanu Jaya > Master Data > Colors.")
    number(doc, "Klik New.")
    number(doc, "Isi Name, misalnya Black Gloss, Candy Red, White Pearl.")
    number(doc, "Isi Code bila digunakan.")
    number(doc, "Klik Save.")

    add_heading(doc, "3.4 Service Master", 2)
    number(doc, "Masuk ke Sentanu Jaya > Master Data > Services.")
    number(doc, "Klik New.")
    number(doc, "Isi Name, misalnya Painting Full Body, Polish, Coating.")
    number(doc, "Pilih Service Type: Raw Preparation, Painting, Polish / Finishing, Rework, atau Other.")
    number(doc, "Jika service punya product terkait untuk accounting/invoice, pilih Related Product.")
    number(doc, "Klik Save.")

    add_heading(doc, "3.5 Defect dan Claim Reason", 2)
    add_para(doc, "Defect dipakai untuk Incoming QC, Final QC, Internal Rework, dan Claim. Claim Reason dipakai sebagai alasan customer melakukan komplain.")
    add_table(
        doc,
        ["Master", "Menu", "Contoh Data"],
        [
            ["Defect", "Sentanu Jaya/Master Data/Defects", "Debu, Belang, Retak, Pecah, Gloss tidak rata"],
            ["Claim Reason", "Sentanu Jaya/Master Data/Claim Reasons", "Warna tidak sesuai, Cat mengelupas, Barang cacat setelah diterima"],
        ],
        widths=[Inches(1.35), Inches(2.4), Inches(2.75)],
    )

    add_heading(doc, "3.6 Customer Pricelist", 2)
    number(doc, "Masuk ke Sentanu Jaya > Master Data > Customer Pricelists.")
    number(doc, "Klik New.")
    number(doc, "Pilih Customer.")
    number(doc, "Pilih Item/Product.")
    number(doc, "Pilih Color dan Service jika harga tergantung warna atau jenis jasa.")
    number(doc, "Isi Selling Price.")
    number(doc, "Sistem menghitung Service Price 10% dan Material Price 90%.")
    number(doc, "Klik Save.")
    add_callout(doc, "Contoh pricing", "Jika Selling Price Rp150.000, sistem menghitung Service Price Rp15.000 dan Material Price Rp135.000.")

    section_break(doc)
    add_heading(doc, "4. Workflow Operasional Barang Masuk sampai Keluar", 1)
    add_heading(doc, "4.1 Tahap 1 - Input Surat Jalan Masuk", 2)
    add_para(doc, "Menu: Sentanu Jaya > Receiving > Surat Jalan Masuk.")
    number(doc, "Klik New.")
    number(doc, "Pilih Customer.")
    number(doc, "Isi Customer SJ Number sesuai nomor surat jalan dari customer.")
    number(doc, "Isi Date Received.")
    number(doc, "Pada tab Items, tambahkan baris item.")
    number(doc, "Pilih Item/Product, Color, Qty, dan UoM.")
    number(doc, "Jika perlu catatan, isi Item Description atau QC Note.")
    number(doc, "Klik Save.")
    number(doc, "Klik tombol Receive untuk mengubah status dari Draft menjadi Received.")
    add_table(
        doc,
        ["Field", "Diisi dengan", "Contoh"],
        [
            ["Customer", "Nama customer pengirim barang", "PT ABC Motor"],
            ["Customer SJ Number", "Nomor SJ dari customer", "SJ-CUST-001"],
            ["Date Received", "Tanggal barang diterima", "2026-06-07"],
            ["Item/Product", "Barang yang diterima", "Roof Garnish Mazda"],
            ["Color", "Warna yang diminta", "Black Gloss"],
            ["Qty", "Jumlah barang", "100"],
            ["QC State", "Status QC line", "Pending, Pass, Reject"],
        ],
        widths=[Inches(1.4), Inches(3.0), Inches(2.1)],
    )

    add_heading(doc, "4.2 Tahap 2 - Incoming QC", 2)
    add_para(doc, "Incoming QC dilakukan dari dokumen Surat Jalan Masuk yang sudah diterima.")
    number(doc, "Buka dokumen Surat Jalan Masuk.")
    number(doc, "Periksa fisik barang: qty sesuai, barang pecah, retak, salah kirim, atau terlalu rusak.")
    number(doc, "Jika semua item lolos, klik QC Pass. Sistem akan menandai line sebagai Pass dan dokumen menjadi QC Passed.")
    number(doc, "Jika barang ditolak, isi Reject Reason di line yang bermasalah, lalu klik QC Reject.")
    number(doc, "Jika barang harus dikembalikan ke customer, klik Return.")
    add_callout(doc, "Aturan operasional", "Barang tidak boleh masuk produksi sebelum Surat Jalan Masuk berstatus QC Passed. Jika status QC Rejected atau Returned, proses berhenti di receiving dan tidak dibuat Production Batch.")

    add_heading(doc, "4.3 Tahap 3 - Membuat Production Batch", 2)
    add_para(doc, "Menu: Sentanu Jaya > Production > Production Batches.")
    number(doc, "Klik New.")
    number(doc, "Pilih SJ Masuk Line dari barang yang sudah QC Passed.")
    number(doc, "Isi Qty batch. Qty batch boleh sebagian dari total qty SJ Masuk Line.")
    number(doc, "Klik Save.")
    number(doc, "Klik Create WO Raw/Paint/Polish.")
    number(doc, "Sistem membuat tiga Work Order: WO Raw, WO Paint, dan WO Polish.")
    add_table(
        doc,
        ["Skenario", "Cara Input Batch"],
        [
            ["Semua barang diproses sekaligus", "Buat 1 Production Batch dengan qty sama dengan qty SJ Masuk Line."],
            ["Barang diproses bertahap", "Buat beberapa Production Batch. Contoh qty 100 dibuat Batch A 50, Batch B 30, Batch C 20."],
            ["Qty batch melebihi SJ Masuk", "Sistem menolak karena total batch tidak boleh lebih besar dari qty line."],
        ],
        widths=[Inches(2.15), Inches(4.35)],
    )

    add_heading(doc, "4.4 Tahap 4 - Menjalankan Work Order", 2)
    add_para(doc, "Menu: Sentanu Jaya > Production > Work Orders.")
    number(doc, "Buka Work Order yang akan dikerjakan.")
    number(doc, "Pilih atau isi PIC jika belum ada.")
    number(doc, "Klik Start saat pekerjaan dimulai.")
    number(doc, "Jika ada pemakaian material, buka tab Material Consumption.")
    number(doc, "Tambahkan material, qty, UoM, dan note.")
    number(doc, "Klik Done saat pekerjaan selesai.")
    add_table(
        doc,
        ["WO Type", "Pekerjaan", "Output"],
        [
            ["WO Raw", "Buffing, Sanding, Putty/Dempul, Epoxy, Primer", "Ready Paint"],
            ["WO Paint", "Base Coat, Color Layer, Clear Coat, Oven", "Ready Finishing"],
            ["WO Polish", "Wet Sanding, Compound, Polishing", "Ready QC"],
            ["WO Rework", "Perbaikan dari internal fail atau customer claim", "Barang siap QC ulang/pengiriman ulang"],
        ],
        widths=[Inches(1.2), Inches(3.2), Inches(2.1)],
    )

    add_heading(doc, "4.5 Tahap 5 - Final QC", 2)
    add_para(doc, "Menu: Sentanu Jaya > Production > Final QC.")
    number(doc, "Klik New.")
    number(doc, "Pilih Production Batch.")
    number(doc, "Isi Date dan Checked By.")
    number(doc, "Pilih Result: Pass atau Fail.")
    number(doc, "Jika Fail, pilih Defect dan isi Note.")
    number(doc, "Klik Save.")
    number(doc, "Klik Apply Result.")
    add_callout(doc, "Hasil Final QC", "Jika Result Pass, batch menjadi QC Passed dan boleh dikirim. Jika Result Fail, batch menjadi QC Failed dan harus masuk Internal Rework sebelum dikirim.")

    add_heading(doc, "4.6 Tahap 6 - Internal Rework jika Final QC Fail", 2)
    add_para(doc, "Menu: Sentanu Jaya > Production > Internal Rework.")
    number(doc, "Klik New.")
    number(doc, "Pilih Production Batch yang gagal QC.")
    number(doc, "Pilih Source Final QC jika tersedia.")
    number(doc, "Pilih Target WO Type: WO Paint atau WO Polish.")
    number(doc, "Pilih Reason/Defect dan isi Note.")
    number(doc, "Klik Save.")
    number(doc, "Klik Start. Sistem membuat Work Order sesuai target dan memulai rework.")
    number(doc, "Setelah rework selesai, klik Done.")
    number(doc, "Lakukan Final QC ulang sampai batch lolos.")

    add_heading(doc, "4.7 Tahap 7 - Surat Jalan Keluar", 2)
    add_para(doc, "Menu: Sentanu Jaya > Delivery > Surat Jalan Keluar.")
    number(doc, "Klik New.")
    number(doc, "Pilih Customer.")
    number(doc, "Isi Date Delivery.")
    number(doc, "Pilih Delivery Type: Normal Delivery.")
    number(doc, "Pada tab Items, tambahkan Production Batch yang akan dikirim.")
    number(doc, "Isi Qty. Untuk partial delivery, qty boleh lebih kecil dari total batch.")
    number(doc, "Klik Save.")
    number(doc, "Klik Ready saat barang siap dikirim.")
    number(doc, "Klik Delivered setelah barang benar-benar dikirim.")
    add_table(
        doc,
        ["Kasus", "Cara Input"],
        [
            ["Kirim seluruh batch", "Isi qty sama dengan qty batch."],
            ["Partial delivery", "Buat SJK pertama untuk qty sebagian, lalu buat SJK berikutnya untuk sisa qty."],
            ["Redelivery claim", "Jangan buat manual sebagai normal delivery. Buat dari Claim Rework agar delivery_type menjadi Redelivery."],
        ],
        widths=[Inches(1.8), Inches(4.7)],
    )

    add_heading(doc, "4.8 Tahap 8 - Invoice dan Payment", 2)
    add_para(doc, "Saat ini integrasi invoice otomatis dari Surat Jalan Keluar belum dibuat. Invoice dilakukan menggunakan modul Accounting standar Odoo dan dapat dihubungkan manual ke Surat Jalan Keluar.")
    number(doc, "Masuk ke Accounting > Customers > Invoices.")
    number(doc, "Klik New.")
    number(doc, "Pilih Customer.")
    number(doc, "Tambahkan invoice line sesuai jasa/barang yang ditagihkan.")
    number(doc, "Gunakan Customer Pricelist sebagai acuan harga.")
    number(doc, "Jika satu invoice menagih beberapa Surat Jalan Keluar, masukkan semua line terkait dalam invoice yang sama.")
    number(doc, "Klik Confirm/Post sesuai workflow Accounting.")
    number(doc, "Setelah customer membayar, klik Register Payment.")
    number(doc, "Kembali ke Sentanu Jaya > Delivery > Surat Jalan Keluar, buka dokumen SJK, pilih Invoice yang sesuai, lalu klik Invoiced.")
    add_callout(doc, "Batasan saat ini", "Satu Invoice dapat menagih beberapa SJK, tetapi relasi invoice ke SJK masih manual. Pengembangan berikutnya perlu membuat wizard Generate Invoice dari beberapa Surat Jalan Keluar.")

    section_break(doc)
    add_heading(doc, "5. Workflow Customer Claim dan Redelivery", 1)
    add_heading(doc, "5.1 Membuat Customer Claim", 2)
    add_para(doc, "Menu: Sentanu Jaya > Claims > Customer Claims.")
    number(doc, "Klik New.")
    number(doc, "Pilih Source SJ Keluar. Claim wajib mengacu ke Surat Jalan Keluar asal.")
    number(doc, "Pilih Claim Reason.")
    number(doc, "Pada tab Claim Items, pilih Delivery Line yang dikomplain.")
    number(doc, "Isi Qty claim, Defect, dan Note.")
    number(doc, "Isi Description jika ada detail komplain dari customer.")
    number(doc, "Klik Save.")
    number(doc, "Klik Review.")

    add_heading(doc, "5.2 Memutuskan Claim Valid atau Invalid", 2)
    add_table(
        doc,
        ["Keputusan", "Tombol", "Dampak"],
        [
            ["Claim tidak valid", "Invalid lalu Close", "Claim ditutup tanpa rework."],
            ["Claim valid", "Valid", "Claim dapat lanjut ke Customer Return dan Create Rework."],
            ["Barang customer sudah kembali", "Customer Return", "Status menjadi Returned."],
        ],
        widths=[Inches(1.7), Inches(1.6), Inches(3.2)],
    )

    add_heading(doc, "5.3 Membuat Claim Rework", 2)
    number(doc, "Buka Customer Claim yang sudah valid.")
    number(doc, "Klik Create Rework.")
    number(doc, "Buka tab Rework atau menu Sentanu Jaya > Claims > Claim Rework.")
    number(doc, "Buka Claim Rework yang dibuat.")
    number(doc, "Klik Start untuk membuat dan menjalankan WO Rework dari item claim.")
    number(doc, "Setelah pekerjaan selesai, klik Done.")

    add_heading(doc, "5.4 Membuat Surat Jalan Pengiriman Ulang", 2)
    number(doc, "Buka Claim Rework yang sudah selesai.")
    number(doc, "Klik Create Redelivery.")
    number(doc, "Sistem membuat Surat Jalan Keluar baru dengan Delivery Type = Redelivery.")
    number(doc, "Nomor redelivery menggunakan sequence SJU, misalnya SJU/2026/00001.")
    number(doc, "Buka redelivery tersebut dari field Redelivery SJ atau dari Sentanu Jaya > Delivery > Surat Jalan Keluar.")
    number(doc, "Klik Ready, lalu Delivered setelah barang pengganti/rework dikirim ulang.")
    add_callout(doc, "Prinsip audit trail claim", "Jangan memakai ulang Surat Jalan Keluar lama untuk hasil rework. Redelivery harus menjadi dokumen baru agar histori SJ awal, claim, rework, dan pengiriman ulang tetap jelas.")

    add_heading(doc, "6. Cabang Keputusan Penting", 1)
    add_table(
        doc,
        ["Kondisi", "Aksi di Sistem", "Status Akhir"],
        [
            ["Incoming QC reject", "Isi Reject Reason, klik QC Reject, lalu Return jika barang dikembalikan.", "SJ Masuk QC Rejected atau Returned"],
            ["Barang lolos incoming", "Klik QC Pass.", "SJ Masuk QC Passed"],
            ["Batch diproses sebagian", "Buat beberapa Production Batch dari line yang sama.", "Beberapa PB dengan total qty <= qty SJ line"],
            ["Final QC fail", "Buat Internal Rework, Start, Done, lalu Final QC ulang.", "Batch kembali siap QC atau QC Passed"],
            ["Partial delivery", "Buat beberapa SJK normal dari batch yang sama dengan qty bertahap.", "Beberapa SJK normal"],
            ["Customer claim invalid", "Review, Invalid, Close.", "Claim Closed"],
            ["Customer claim valid", "Valid, Customer Return, Create Rework, Start, Done, Create Redelivery.", "Claim Done dan SJU dibuat"],
        ],
        widths=[Inches(2.0), Inches(3.25), Inches(1.25)],
    )

    add_heading(doc, "7. Checklist Harian per Bagian", 1)
    add_table(
        doc,
        ["Bagian", "Checklist Harian"],
        [
            ["Admin Receiving", "Input semua SJ Masuk, cek nomor SJ customer, pastikan line item/qty/warna benar, jalankan Receive dan Incoming QC."],
            ["QC Incoming", "Periksa fisik barang, tandai pass/reject, isi defect/reject reason untuk barang bermasalah."],
            ["Supervisor Produksi", "Buat Production Batch, pastikan qty batch tidak melebihi SJ line, buat WO Raw/Paint/Polish."],
            ["PIC Produksi", "Start WO saat mulai kerja, input material consumption, Done saat selesai."],
            ["QC Final", "Buat Final QC, isi result, defect, note, dan Apply Result."],
            ["Delivery", "Buat SJK dari batch QC Passed, gunakan partial delivery jika perlu, klik Ready dan Delivered."],
            ["Accounting", "Buat invoice dari SJK yang sudah delivered, register payment, hubungkan invoice ke SJK, klik Invoiced."],
            ["Admin Claim", "Input claim dari SJK asal, review valid/invalid, buat rework dan redelivery jika valid."],
        ],
        widths=[Inches(1.55), Inches(4.95)],
    )

    add_heading(doc, "8. Menu Referensi Cepat", 1)
    add_table(
        doc,
        ["Kebutuhan", "Menu"],
        [
            ["Membuat customer", "Contacts"],
            ["Mengisi tax customer", "Contacts > buka customer > Tax Information"],
            ["Membuat user", "Settings > Users & Companies > Users"],
            ["Membuat employee/PIC", "Employees"],
            ["Membuat gudang", "Inventory > Configuration > Warehouses"],
            ["Membuat lokasi gudang", "Inventory > Configuration > Locations"],
            ["Master warna", "Sentanu Jaya > Master Data > Colors"],
            ["Master jasa", "Sentanu Jaya > Master Data > Services"],
            ["Master defect", "Sentanu Jaya > Master Data > Defects"],
            ["Master alasan claim", "Sentanu Jaya > Master Data > Claim Reasons"],
            ["Customer pricelist", "Sentanu Jaya > Master Data > Customer Pricelists"],
            ["Tax profile", "Sentanu Jaya > Configuration > Tax Profiles"],
            ["Surat Jalan Masuk", "Sentanu Jaya > Receiving > Surat Jalan Masuk"],
            ["Production Batch", "Sentanu Jaya > Production > Production Batches"],
            ["Work Order", "Sentanu Jaya > Production > Work Orders"],
            ["Final QC", "Sentanu Jaya > Production > Final QC"],
            ["Internal Rework", "Sentanu Jaya > Production > Internal Rework"],
            ["Surat Jalan Keluar", "Sentanu Jaya > Delivery > Surat Jalan Keluar"],
            ["Customer Claim", "Sentanu Jaya > Claims > Customer Claims"],
            ["Claim Rework", "Sentanu Jaya > Claims > Claim Rework"],
            ["Invoice", "Accounting > Customers > Invoices"],
            ["Payment", "Accounting > Customers > Payments atau dari Register Payment pada Invoice"],
        ],
        widths=[Inches(2.0), Inches(4.5)],
    )

    add_heading(doc, "9. Troubleshooting Operasional", 1)
    add_table(
        doc,
        ["Masalah", "Kemungkinan Penyebab", "Solusi"],
        [
            ["Menu Sentanu Jaya tidak muncul", "User belum punya akses internal atau browser belum refresh registry.", "Logout/login kembali. Jika tetap tidak muncul, admin cek akses user di Settings > Users."],
            ["Tidak bisa memilih customer", "Customer belum dibuat di Contacts.", "Buat customer terlebih dahulu dari Contacts."],
            ["Tidak bisa memilih item", "Product belum dibuat.", "Buat item dari Inventory/Sales > Products."],
            ["Production Batch ditolak", "Total qty batch melebihi qty SJ Masuk Line.", "Cek batch yang sudah dibuat dari line tersebut. Buat qty batch yang tidak melebihi sisa qty."],
            ["Delivery line ditolak", "Qty normal delivery melebihi qty batch.", "Gunakan qty sesuai sisa batch. Untuk pengiriman ulang claim, buat dari Claim Rework agar menjadi Redelivery."],
            ["Claim line tidak muncul", "Source SJ Keluar belum dipilih atau delivery line tidak berasal dari SJ tersebut.", "Pilih Source SJ Keluar dulu, lalu pilih delivery line yang sesuai."],
            ["Invoice belum terhubung ke SJK", "Integrasi invoice otomatis belum dibuat.", "Buat invoice manual di Accounting, lalu pilih invoice tersebut di dokumen SJK dan klik Invoiced."],
        ],
        widths=[Inches(1.8), Inches(2.25), Inches(2.45)],
    )

    add_heading(doc, "10. Contoh Simulasi Transaksi End-to-End", 1)
    add_para(doc, "Contoh ini dapat dipakai saat training user untuk mencoba alur dari awal sampai akhir. Data contoh bisa disesuaikan dengan data real Sentanu Jaya.")
    add_table(
        doc,
        ["Data", "Nilai Contoh"],
        [
            ["Customer", "PT ABC Motor"],
            ["Customer SJ Number", "SJ-ABC-001"],
            ["Item", "Roof Garnish Mazda"],
            ["Warna", "Black Gloss"],
            ["Qty Masuk", "100 pcs"],
            ["Harga Jual", "Rp150.000 per pcs"],
            ["Batch Produksi", "Batch A 50 pcs, Batch B 30 pcs, Batch C 20 pcs"],
            ["Delivery", "SJK-001 50 pcs, SJK-002 50 pcs"],
            ["Claim", "2 pcs warna tidak sesuai dari SJK-001"],
            ["Redelivery", "SJU-001 2 pcs setelah claim rework selesai"],
        ],
        widths=[Inches(1.7), Inches(4.8)],
    )
    add_heading(doc, "10.1 Langkah Simulasi", 2)
    number(doc, "Buat customer PT ABC Motor di Contacts dan isi Tax Information.")
    number(doc, "Buat product Roof Garnish Mazda.")
    number(doc, "Buat color Black Gloss.")
    number(doc, "Buat service Painting dan customer pricelist Rp150.000.")
    number(doc, "Buat Surat Jalan Masuk dari PT ABC Motor dengan nomor SJ-ABC-001 dan qty 100 pcs.")
    number(doc, "Klik Receive, lalu QC Pass.")
    number(doc, "Buat Production Batch A qty 50, Batch B qty 30, dan Batch C qty 20 dari line SJ Masuk tersebut.")
    number(doc, "Pada setiap batch, klik Create WO Raw/Paint/Polish.")
    number(doc, "Jalankan setiap WO: Start, isi material consumption jika ada, lalu Done.")
    number(doc, "Buat Final QC untuk setiap batch. Jika Pass, klik Apply Result.")
    number(doc, "Buat SJK normal pertama qty 50 dari Batch A, klik Ready lalu Delivered.")
    number(doc, "Buat SJK normal kedua untuk sisa barang sesuai batch yang sudah siap, klik Ready lalu Delivered.")
    number(doc, "Buat invoice manual di Accounting untuk SJK yang sudah delivered.")
    number(doc, "Jika ada claim 2 pcs dari SJK pertama, buat Customer Claim dengan Source SJ Keluar SJK pertama.")
    number(doc, "Isi claim line 2 pcs, pilih defect, klik Review, Valid, Customer Return, lalu Create Rework.")
    number(doc, "Buka Claim Rework, klik Start, setelah selesai klik Done.")
    number(doc, "Klik Create Redelivery, buka SJU yang dibuat, klik Ready lalu Delivered.")

    add_heading(doc, "11. Catatan Batasan Versi Saat Ini", 1)
    bullet(doc, "Invoice dari Surat Jalan Keluar belum otomatis; accounting masih membuat invoice manual di modul Accounting.")
    bullet(doc, "Stock movement dan pemakaian inventory material belum otomatis mengurangi stok; material consumption masih berupa pencatatan operasional.")
    bullet(doc, "Report PDF untuk SJM, SJK, SJU, claim, dan WO belum dibuat.")
    bullet(doc, "Security group masih umum untuk internal user; perlu dibuat role detail per bagian.")
    bullet(doc, "Attachment foto QC dan claim belum dibuat khusus di form, tetapi dapat dikembangkan menggunakan attachment/chatter.")

    add_heading(doc, "12. Rekomendasi Penggunaan Saat Go-Live", 1)
    number(doc, "Isi master data dulu: customer, item, warna, jasa, defect, claim reason, tax profile, dan customer pricelist.")
    number(doc, "Gunakan satu dokumen Surat Jalan Masuk untuk setiap surat jalan dari customer.")
    number(doc, "Jangan membuat Production Batch untuk barang yang belum QC Passed.")
    number(doc, "Gunakan Production Batch jika barang dari satu SJ perlu diproses bertahap.")
    number(doc, "Selalu klik Start dan Done pada WO agar histori waktu produksi terbaca.")
    number(doc, "Final QC wajib dilakukan sebelum membuat Surat Jalan Keluar.")
    number(doc, "Customer Claim harus selalu mengacu ke SJK asal.")
    number(doc, "Redelivery harus dibuat dari Claim Rework, bukan dari SJK lama.")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "Panduan Operasional ERP Sentanu Jaya - Odoo 17 Community"
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("666666")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
